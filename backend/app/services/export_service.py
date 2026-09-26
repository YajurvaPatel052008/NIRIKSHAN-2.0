import tempfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

from app.services.validation_service import (
    check_version_status,
    get_certification_guidance,
)


def _clean_text(value: Any, fallback: str = "—") -> str:
    if value is None:
        return fallback
    text = str(value).strip()
    if not text or text.lower() in {"null", "none"}:
        return fallback
    return text


def _analysis_summary(analysis: Any) -> dict[str, str]:
    requirements = analysis.extracted_requirements or {}
    product = _clean_text(requirements.get("product"), "")
    intended_use = _clean_text(requirements.get("intended_use"), "")
    summary_parts = []
    if product:
        summary_parts.append(f"Product: {product}.")
    if intended_use:
        summary_parts.append(f"Intended use: {intended_use}.")
    return {
        "title": _clean_text(analysis.title, "Untitled analysis"),
        "product": product or "Not specified",
        "intended_use": intended_use or "Not specified",
        "product_summary": " ".join(summary_parts) or "Product and intended use were not specified.",
    }


def _certification_requirement(standard: Any) -> str:
    guidance = get_certification_guidance(standard)
    schemes = guidance.get("applicable_schemes", [])
    if not schemes:
        return "No certification metadata available"
    return "; ".join(
        f"{'Mandatory' if scheme.get('mandatory') else 'Optional'}: "
        f"{_clean_text(scheme.get('scheme_name'), 'Certification scheme')}"
        for scheme in schemes
    )


def _recommendation_content(
    accepted_recommendations: list[Any],
) -> list[dict[str, str]]:
    content = []
    for recommendation in accepted_recommendations:
        standard = recommendation.standard
        version = check_version_status(standard)
        relationship_type = (
            recommendation.relationship_type.value
            if recommendation.relationship_type
            else None
        )
        amendments = version.get("amendments") or []
        version_parts = []
        if standard.version_edition:
            version_parts.append(f"Edition: {standard.version_edition}")
        if amendments:
            version_parts.append(f"Amendments: {', '.join(map(str, amendments))}")
        version_parts.append("Verify current status with BIS")
        content.append(
            {
                "is_number": _clean_text(standard.is_number),
                "title": _clean_text(standard.title),
                "relationship": (
                    f"Related — {relationship_type.replace('_', ' ')}"
                    if relationship_type
                    else "Primary"
                ),
                "relationship_type": relationship_type or "",
                "version_amendment_note": "; ".join(version_parts),
                "certification_requirement": _certification_requirement(standard),
                "annotation": _clean_text(recommendation.reviewer_notes, ""),
            }
        )
    return content


def _apply_overrides(
    summary: dict[str, str],
    standards: list[dict[str, str]],
    overrides: dict[str, Any] | None,
) -> tuple[dict[str, str], list[dict[str, str]]]:
    if not overrides:
        return summary, standards

    edited_summary = dict(summary)
    product_summary = overrides.get("product_summary")
    if product_summary is not None:
        edited_summary["product_summary"] = _clean_text(product_summary, "")

    changes_by_number = {
        item["is_number"]: item
        for item in overrides.get("standards", [])
        if isinstance(item, dict) and isinstance(item.get("is_number"), str)
    }
    edited_standards = []
    for standard in standards:
        edited = dict(standard)
        changes = changes_by_number.get(standard["is_number"], {})
        for field in (
            "title",
            "relationship",
            "version_amendment_note",
            "certification_requirement",
            "annotation",
        ):
            if field in changes and changes[field] is not None:
                edited[field] = _clean_text(changes[field], "")
        edited_standards.append(edited)
    return edited_summary, edited_standards


def _specification_content(
    analysis: Any,
    accepted_recommendations: list[Any],
    overrides: dict[str, Any] | None = None,
) -> dict[str, Any]:
    summary, standards = _apply_overrides(
        _analysis_summary(analysis),
        _recommendation_content(accepted_recommendations),
        overrides,
    )
    return {
        "document_type": "tender-ready specification",
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "analysis": summary,
        "accepted_standards": standards,
    }


def generate_specification_json(
    analysis: Any,
    accepted_recommendations: list[Any],
    overrides: dict[str, Any] | None = None,
) -> dict[str, Any]:
    return _specification_content(analysis, accepted_recommendations, overrides)


def generate_specification_docx(
    analysis: Any,
    accepted_recommendations: list[Any],
    overrides: dict[str, Any] | None = None,
) -> str:
    specification = _specification_content(
        analysis, accepted_recommendations, overrides
    )
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)

    title = document.add_heading("Tender-Ready Specification", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title_run = title.runs[0]
    title_run.font.color.rgb = RGBColor(11, 61, 110)
    title_run.font.size = Pt(22)

    document.add_paragraph(f"Analysis: {specification['analysis']['title']}")
    generated_at = datetime.fromisoformat(
        specification["generated_at"].replace("Z", "+00:00")
    )
    document.add_paragraph(f"Generated: {generated_at.astimezone().strftime('%d %B %Y')}")

    document.add_heading("Product & Use Case", level=2)
    document.add_paragraph(specification["analysis"]["product_summary"])

    document.add_heading("Recommended Standards", level=2)
    table = document.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    headers = [
        "IS Number",
        "Title",
        "Relationship",
        "Version/Amendment",
        "Certification",
    ]
    for cell, header in zip(table.rows[0].cells, headers):
        cell.text = header
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True
                run.font.color.rgb = RGBColor(11, 61, 110)

    for standard in specification["accepted_standards"]:
        cells = table.add_row().cells
        values = [
            standard["is_number"],
            standard["title"],
            standard["relationship"] or "—",
            standard["version_amendment_note"] or "—",
            standard["certification_requirement"] or "—",
        ]
        for cell, value in zip(cells, values):
            cell.text = _clean_text(value)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        if standard["annotation"]:
            note_row = table.add_row()
            note_cell = note_row.cells[0]
            for cell in note_row.cells[1:]:
                note_cell = note_cell.merge(cell)
            paragraph = note_row.cells[0].paragraphs[0]
            note_label = paragraph.add_run("Officer note: ")
            note_label.bold = True
            paragraph.add_run(standard["annotation"])

    for paragraph in document.paragraphs:
        paragraph.paragraph_format.space_after = Pt(6)
    for table_row in table.rows:
        for cell in table_row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)

    output = tempfile.NamedTemporaryFile(
        prefix="NIRIKSHAN-specification-",
        suffix=".docx",
        delete=False,
    )
    output.close()
    path = Path(output.name)
    document.save(path)
    return str(path)
